"""Control externo: compara el cálculo Python con la planilla entregada.

La planilla no participa en el cálculo ni en la exportación a Unity.
"""
from pathlib import Path
import copy
import sys
import json
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'analysis' / 'capacity'))
from capacidad import interaction_points


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / 'data' / 'reinforcement' / 'Comprobacion diagrama de interaccion corregida.xlsx'


def main():
    active = json.loads((ROOT / 'data' / 'parameters' / 'parametros.json').read_text(encoding='utf-8'))['columna']
    sheet = openpyxl.load_workbook(BOOK, data_only=True, read_only=True)['Pregunta 2']
    # The external spreadsheet specifies five Ø22 bars per face; the active
    # building uses Ø28. A benchmark is meaningful only at identical inputs.
    # This comparison does not substitute the active model's P-M calculation.
    cfg = copy.deepcopy(active)
    cfg.update(b_m=sheet['B3'].value / 1000, h_m=sheet['B4'].value / 1000,
               fc_MPa=sheet['B7'].value, fy_MPa=sheet['D3'].value,
               Es_MPa=sheet['D4'].value * 1000,
               beta1_interaccion=sheet['B9'].value,
               eps_cu_interaccion=sheet['B10'].value,
               barras_por_cara=int(sheet['D6'].value),
               recubrimiento_al_centro_barra_m=sheet['D16'].value / 1000,
               factor_compresion_max=sheet['H26'].value / sheet['H25'].value)
    cfg['diametros_por_cara_m'] = [sheet.cell(row, 4).value / 1000 for row in range(11, 16)]
    cfg['diametro_m'] = cfg['diametros_por_cara_m'][0]
    calculated = interaction_points(cfg)
    reference = [dict(punto=sheet.cell(row, 6).value,
                      P_kN=float(sheet.cell(row, 7).value),
                      M_kNm=float(sheet.cell(row, 8).value))
                 for row in range(3, 10)]
    max_dp = max(abs(a['P_kN']-b['P_kN']) for a, b in zip(calculated, reference))
    max_dm = max(abs(a['M_kNm']-b['M_kNm']) for a, b in zip(calculated, reference))
    print(f'Diferencia máxima P: {max_dp:.12g} kN')
    print(f'Diferencia máxima M: {max_dm:.12g} kN·m')
    print('Armadura del control externo: ' + ', '.join(f'{d*1000:g}' for d in cfg['diametros_por_cara_m'])
          + ' mm; armadura activa: ' + ', '.join(f'{d*1000:g}' for d in active['diametros_por_cara_m']) + ' mm.')
    if max_dp > 1e-8 or max_dm > 1e-8:
        raise SystemExit('REVISAR: el cálculo independiente no coincide con el control externo')
    print('Interacción P-M con las entradas de la planilla: OK (no valida la armadura activa)')


if __name__ == '__main__':
    main()
