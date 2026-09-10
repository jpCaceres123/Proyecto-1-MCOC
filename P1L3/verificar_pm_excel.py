"""Control externo: compara el cálculo Python con la planilla entregada.

La planilla no participa en el cálculo ni en la exportación a Unity.
"""
from pathlib import Path
import json
import openpyxl

from capacidad import interaction_points


ROOT = Path(__file__).resolve().parent
BOOK = ROOT.parent / 'Comprobacion diagrama de interaccion corregida.xlsx'


def main():
    cfg = json.loads((ROOT / 'parametros.json').read_text(encoding='utf-8'))['columna']
    calculated = interaction_points(cfg)
    sheet = openpyxl.load_workbook(BOOK, data_only=True, read_only=True)['Pregunta 2']
    reference = [dict(punto=sheet.cell(row, 6).value,
                      P_kN=float(sheet.cell(row, 7).value),
                      M_kNm=float(sheet.cell(row, 8).value))
                 for row in range(3, 10)]
    max_dp = max(abs(a['P_kN']-b['P_kN']) for a, b in zip(calculated, reference))
    max_dm = max(abs(a['M_kNm']-b['M_kNm']) for a, b in zip(calculated, reference))
    print(f'Diferencia máxima P: {max_dp:.12g} kN')
    print(f'Diferencia máxima M: {max_dm:.12g} kN·m')
    if max_dp > 1e-8 or max_dm > 1e-8:
        raise SystemExit('REVISAR: el cálculo independiente no coincide con el control externo')
    print('Interacción P-M: OK')


if __name__ == '__main__':
    main()
